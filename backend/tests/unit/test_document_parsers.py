"""Unit tests for document parsers (PDF, DOCX, TXT) and DocumentParserFactory."""

import os
import tempfile
from pathlib import Path

import pytest
import docx

from app.providers.document.factory import DocumentParserFactory, parser_factory
from app.providers.document.pdf_parser import PDFParser
from app.providers.document.docx_parser import DOCXParser
from app.providers.document.txt_parser import TXTParser


# === Factory Tests ===

class TestDocumentParserFactory:
    def test_get_parser_pdf(self):
        factory = DocumentParserFactory()
        parser = factory.get_parser(".pdf")
        assert isinstance(parser, PDFParser)

    def test_get_parser_filename_pdf(self):
        factory = DocumentParserFactory()
        parser = factory.get_parser("resume.PDF")
        assert isinstance(parser, PDFParser)

    def test_get_parser_docx(self):
        factory = DocumentParserFactory()
        parser = factory.get_parser("job_desc.docx")
        assert isinstance(parser, DOCXParser)

    def test_get_parser_txt(self):
        factory = DocumentParserFactory()
        parser = factory.get_parser("notes.txt")
        assert isinstance(parser, TXTParser)

    def test_get_parser_md(self):
        factory = DocumentParserFactory()
        parser = factory.get_parser("README.md")
        assert isinstance(parser, TXTParser)

    def test_unsupported_format_raises_value_error(self):
        factory = DocumentParserFactory()
        with pytest.raises(ValueError, match="chưa được hỗ trợ"):
            factory.get_parser("spreadsheet.xlsx")


# === TXT Parser Tests ===

class TestTXTParser:
    @pytest.mark.asyncio
    async def test_extract_utf8_text(self, tmp_path: Path):
        file_path = tmp_path / "sample.txt"
        content = "Họ và tên: Võ Văn Tuấn\nKỹ năng: C++, Qt, Linux\n\nDự án: Xe tự hành"
        file_path.write_text(content, encoding="utf-8")

        parser = TXTParser()
        extracted = await parser.extract_text(str(file_path))
        assert "Võ Văn Tuấn" in extracted
        assert "C++, Qt, Linux" in extracted
        assert "Dự án: Xe tự hành" in extracted

    @pytest.mark.asyncio
    async def test_extract_empty_file_raises_value_error(self, tmp_path: Path):
        empty_file = tmp_path / "empty.txt"
        empty_file.write_text("", encoding="utf-8")

        parser = TXTParser()
        with pytest.raises(ValueError, match="(?i)(empty|trống)"):
            await parser.extract_text(str(empty_file))

    @pytest.mark.asyncio
    async def test_nonexistent_file_raises_not_found(self):
        parser = TXTParser()
        with pytest.raises(FileNotFoundError):
            await parser.extract_text("non_existent_file_12345.txt")


# === DOCX Parser Tests ===

class TestDOCXParser:
    @pytest.mark.asyncio
    async def test_extract_docx_paragraphs_and_table(self, tmp_path: Path):
        file_path = tmp_path / "test.docx"
        doc = docx.Document()
        doc.add_heading("CV - Võ Văn Tuấn", level=1)
        doc.add_paragraph("Kỹ sư phần mềm C++ Fresher.")

        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "Kỹ năng"
        table.rows[0].cells[1].text = "Mức độ"
        table.rows[1].cells[0].text = "C++ / Qt"
        table.rows[1].cells[1].text = "Thành thạo"

        doc.save(str(file_path))

        parser = DOCXParser()
        extracted = await parser.extract_text(str(file_path))
        assert "CV - Võ Văn Tuấn" in extracted
        assert "Kỹ sư phần mềm C++ Fresher." in extracted
        assert "C++ / Qt" in extracted
        assert "Thành thạo" in extracted

    @pytest.mark.asyncio
    async def test_corrupted_docx_raises_value_error(self, tmp_path: Path):
        bad_file = tmp_path / "corrupted.docx"
        bad_file.write_bytes(b"not a valid zip or docx content")

        parser = DOCXParser()
        with pytest.raises(ValueError, match="Không thể mở file DOCX"):
            await parser.extract_text(str(bad_file))


# === PDF Parser Tests ===

class TestPDFParser:
    @pytest.mark.asyncio
    async def test_extract_real_pdf_cv(self):
        # Test with the existing real PDF in the workspace
        real_pdf = r"D:\CV\Automative\CV_Vo_Van_Tuan_C++_Automative.pdf"
        if not os.path.exists(real_pdf):
            pytest.skip(f"Test file not found: {real_pdf}")

        parser = PDFParser()
        extracted = await parser.extract_text(real_pdf)
        assert len(extracted) > 100
        # Check that C++ or relevant automotive content is captured
        assert "C++" in extracted or "TUAN" in extracted.upper()

    @pytest.mark.asyncio
    async def test_corrupted_pdf_raises_value_error(self, tmp_path: Path):
        bad_pdf = tmp_path / "bad.pdf"
        bad_pdf.write_bytes(b"this is not a valid pdf header")

        parser = PDFParser()
        with pytest.raises(ValueError, match="Không thể mở file PDF"):
            await parser.extract_text(str(bad_pdf))
