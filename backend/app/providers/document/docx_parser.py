"""DOCX document parser using python-docx."""

import docx
import structlog

from app.providers.document.base import DocumentParser
from app.utils.text_utils import clean_whitespace

logger = structlog.get_logger()


class DOCXParser(DocumentParser):
    """Extracts text content from DOCX files."""

    SUPPORTED_EXTENSIONS = {".docx"}

    def supports(self, file_extension: str) -> bool:
        return file_extension.lower() in self.SUPPORTED_EXTENSIONS

    async def extract_text(self, file_path: str) -> str:
        """Extract text from DOCX document including paragraphs and tables.

        Args:
            file_path: Path to the DOCX file.

        Returns:
            Extracted and normalized text.

        Raises:
            FileNotFoundError: If file does not exist.
            ValueError: If file is corrupted or contains no readable text.
        """
        path = self.validate_file_exists(file_path)

        try:
            doc = docx.Document(str(path))
        except Exception as e:
            logger.error("docx_open_failed", path=str(path), error=str(e))
            raise ValueError(f"Không thể mở file DOCX (file hỏng hoặc định dạng sai): {e}") from e

        text_parts: list[str] = []

        # 1. Paragraphs
        for paragraph in doc.paragraphs:
            p_text = paragraph.text.strip()
            if p_text:
                text_parts.append(p_text)

        # 2. Tables
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    # Remove consecutive duplicate cell texts from merged cells
                    unique_cells: list[str] = []
                    for c in row_cells:
                        if not unique_cells or unique_cells[-1] != c:
                            unique_cells.append(c)
                    text_parts.append(" | ".join(unique_cells))

        full_text = "\n\n".join(text_parts)
        cleaned_text = clean_whitespace(full_text)

        if not cleaned_text.strip():
            raise ValueError("File DOCX không chứa nội dung văn bản.")

        logger.info(
            "docx_extracted_successfully",
            path=str(path),
            paragraphs=len(doc.paragraphs),
            tables=len(doc.tables),
            char_count=len(cleaned_text),
        )
        return cleaned_text
